"""Document parser worker. Outputs bounded segments with explicit source locations."""
from __future__ import annotations
import csv,io,json,re,sys,zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

MAX_TEXT=2_000_000
TEXT_SUFFIXES={'.txt','.md','.csv','.json','.jsonl','.py','.js','.ts','.tsx','.jsx','.html','.css','.yaml','.yml','.xml','.sql','.log','.sh','.toml','.ini','.ipynb','.go','.rs','.java','.c','.cpp','.h','.vue','.svelte','.env.example'}
IMAGE_SUFFIXES={'.png','.jpg','.jpeg','.webp','.gif','.bmp','.tiff','.tif'}
SUPPORTED=TEXT_SUFFIXES|IMAGE_SUFFIXES|{'.pdf','.docx','.xlsx','.pptx','.zip'}

def decode(data):
    for encoding in (['utf-16'] if data.startswith((b'\xff\xfe',b'\xfe\xff')) else ['utf-8-sig','gb18030']):
        try:return data.decode(encoding)
        except UnicodeError:pass
    return data.decode('utf-8',errors='replace')

def checked_zip(path):
    archive=zipfile.ZipFile(path)
    entries=archive.infolist()
    if len(entries)>3000 or sum(x.file_size for x in entries)>100_000_000:
        archive.close();raise ValueError('Archive exceeds entry or expanded-size limits')
    for entry in entries:
        name=Path(entry.filename.replace('\\','/'))
        if name.is_absolute() or '..' in name.parts or (entry.external_attr>>16)&0o170000==0o120000:
            archive.close();raise ValueError('Unsafe archive entry')
        if entry.flag_bits&1:archive.close();raise ValueError('Password-protected archives are not supported')
    return archive

def parse(path,name):
    path=Path(path);suffix=Path(name).suffix.lower();segments=[];warnings=[];total=0
    def add(location,text):
        nonlocal total
        text=str(text)
        if not text.strip():return
        if total>=MAX_TEXT:return
        text=text[:MAX_TEXT-total];total+=len(text)
        for offset in range(0,len(text),6000):segments.append({'location':location,'text':text[offset:offset+6000]})
    if suffix in IMAGE_SUFFIXES:
        from PIL import Image
        Image.MAX_IMAGE_PIXELS=25_000_000
        with Image.open(path) as image:
            if image.width*image.height>25_000_000:raise ValueError("Image exceeds 25 million pixels")
            image.verify()
        with Image.open(path) as image:
            return {'segments':[],'warnings':[],'kind':'image','width':image.width,'height':image.height,'format':image.format}
    if suffix in TEXT_SUFFIXES:
        text=decode(path.read_bytes())
        lines=text.splitlines()
        for index in range(0,len(lines),60):add(f'lines {index+1}-{min(index+60,len(lines))}','\n'.join(lines[index:index+60]))
    elif suffix=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(path)
        if reader.is_encrypted:raise ValueError('Password-protected PDF requires an unlocked copy')
        if len(reader.pages)>200:raise ValueError('PDF exceeds the 200-page limit')
        for index,page in enumerate(reader.pages):
            text=page.extract_text() or ''
            if text.strip():add(f'page {index+1}',text)
            else:
                warnings.append(f'Page {index+1} requires image_understand(file_id, page={index+1}) for OCR')
                add(f'page {index+1}','[No extractable text. Read this page with image_understand.]')
    elif suffix in {'.docx','.pptx'}:
        with checked_zip(path) as archive:
            if suffix=='.docx':
                root=ET.fromstring(archive.read('word/document.xml'))
                ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                for index,paragraph in enumerate(root.findall('.//w:p',ns)):
                    add(f'paragraph {index+1}',''.join(x.text or '' for x in paragraph.findall('.//w:t',ns)))
            else:
                slides=sorted((n for n in archive.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)),key=lambda n:int(re.search(r'slide(\d+)',n).group(1)))
                ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
                for index,name in enumerate(slides):
                    root=ET.fromstring(archive.read(name));add(f'slide {index+1}','\n'.join(x.text or '' for x in root.findall('.//a:t',ns)))
    elif suffix=='.xlsx':
        with checked_zip(path):pass
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(path.read_bytes()),read_only=True,data_only=False)
        try:
            for sheet in book:
                for i,row in enumerate(sheet.iter_rows(max_col=200)):
                    if i>=10000:warnings.append(f'{sheet.title}: only first 10000 rows parsed');break
                    text=' | '.join(f'{cell.coordinate}={cell.value}' for cell in row[:200] if cell.value is not None)
                    add(f'sheet {sheet.title}, row {i+1}',text)
            warnings.append('Formulas are preserved, not recalculated. Use table_calculate for numeric aggregates.')
        finally:book.close()
    elif suffix=='.zip':
        with checked_zip(path) as archive:
            for item in archive.infolist():
                if item.is_dir():continue
                if Path(item.filename).suffix.lower() in TEXT_SUFFIXES and item.file_size<=1_000_000:
                    add(f'archive entry {item.filename}',decode(archive.read(item)))
                else:add(f'archive entry {item.filename}',f'[Binary or large file; {item.file_size} bytes. Extract explicitly in the workspace if needed.]')
    else:raise ValueError('Unsupported file format')
    if suffix in {'.docx','.pptx','.xlsx'}:
        with checked_zip(path) as archive:
            images=sorted(name for name in archive.namelist() if '/media/' in name and Path(name).suffix.lower() in IMAGE_SUFFIXES)
            for index,name in enumerate(images,1):add(f'embedded image {index}',f'{name} ? use file_extract(file_id, kind="image", index={index}) to inspect this image')
    if total>=MAX_TEXT:warnings.append('Parsed text was truncated at 2 million characters; split the document for complete analysis')
    return {'kind':'document','segments':segments,'warnings':warnings,'characters':total}

if __name__=='__main__':
    if sys.platform!='win32':
        import resource
        resource.setrlimit(resource.RLIMIT_CPU,(30,30))
        resource.setrlimit(resource.RLIMIT_AS,(268435456,268435456))
    source,name,destination=sys.argv[1:4]
    try:result=parse(source,name)
    except Exception as exc:result={'error':f'{type(exc).__name__}: {exc}'}
    Path(destination).write_text(json.dumps(result,ensure_ascii=False),encoding='utf-8')
