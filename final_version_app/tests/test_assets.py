"""File parsing, durable uploads, source citations and numeric correctness."""
import io,json,time,zipfile
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from final_version_app.document_parser import parse
from final_version_app.storage.assets import AssetStore
from final_version_app.domain.tables import calculate_table
from final_version_app.tests.test_workspace_auth import make_app,CREDENTIALS

def wait_ready(store,asset):
    for _ in range(100):
        metadata=store.metadata(asset['id'])
        if metadata['status'] in {'ready','failed'}:break
        time.sleep(.05)
    assert metadata['status']=='ready',metadata
    return metadata

def zip_fixture(entries):
    data=io.BytesIO()
    with zipfile.ZipFile(data,'w',zipfile.ZIP_DEFLATED) as archive:
        for name,text in entries.items():archive.writestr(name,text)
    return data.getvalue()

def test_upload_auth_persistence_and_reference_prompt(tmp_path,monkeypatch):
    store=AssetStore(tmp_path)
    monkeypatch.setattr('final_version_app.storage.assets.get_asset_store',lambda:store)
    client=TestClient(make_app(tmp_path))
    assert client.post('/api/files?name=budget.csv',content=b'amount\n12\n').status_code==401
    client.post('/api/auth/setup',json=CREDENTIALS)
    uploaded=client.post('/api/files?name=budget.csv',content=b'category,amount\nA,12.10\nB,3.20\nA,2.90\n')
    assert uploaded.status_code==201
    asset=uploaded.json();wait_ready(store,asset)
    assert client.get(asset['download_url']).content.startswith(b'category,amount')
    parsed=store.read(asset['id']);assert parsed['segments'][0]['citation'].startswith('budget.csv')
    result=calculate_table(store,asset['id'],'amount','sum',filter_column='category',filter_value='A')
    assert result['value']=='15.00' and result['matched_rows']==2
    from final_version_app.api import _prompt_with_attachments,AttachmentPayload
    prompt=_prompt_with_attachments('Summarize',[AttachmentPayload(name='budget.csv',file_id=asset['id'])])
    assert asset['id'] in prompt and '12.10' not in prompt
    rebuilt=AssetStore(tmp_path);assert rebuilt.metadata(asset['id'])['status']=='ready'
    client.post('/api/auth/logout');assert client.get(asset['download_url']).status_code==401
    store.executor.shutdown();rebuilt.executor.shutdown()

def test_office_pdf_image_and_zip_sources(tmp_path):
    from PIL import Image
    from pypdf import PdfWriter
    from pypdf.generic import NameObject,DictionaryObject,DecodedStreamObject
    from openpyxl import Workbook
    doc=tmp_path/'sample.docx';doc.write_bytes(zip_fixture({'word/document.xml':'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Revenue target is 480.</w:t></w:r></w:p></w:body></w:document>'}))
    assert parse(doc,doc.name)['segments'][0]['location']=='paragraph 1'
    deck=tmp_path/'sample.pptx';deck.write_bytes(zip_fixture({'ppt/slides/slide1.xml':'<p:sld xmlns:p="p" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:t>Launch in October</a:t></p:sld>'}))
    assert parse(deck,deck.name)['segments'][0]['text']=='Launch in October'
    sheet=tmp_path/'sample.xlsx';book=Workbook();book.active.append(['Amount']);book.active.append([120]);book.save(sheet)
    assert 'A2=120' in str(parse(sheet,sheet.name)['segments'])
    pdf=tmp_path/'sample.pdf';writer=PdfWriter();page=writer.add_blank_page(width=300,height=300)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 14 Tf 30 250 Td (Revenue 480) Tj ET');page[NameObject('/Contents')]=writer._add_object(stream)
    writer.add_blank_page(width=300,height=300);writer.write(pdf)
    result=parse(pdf,pdf.name)
    assert result['segments'][0]['location']=='page 1' and 'Revenue 480' in result['segments'][0]['text']
    assert 'OCR' in result['warnings'][0]
    png=tmp_path/'chart.png';Image.new('RGB',(320,200),'white').save(png)
    assert parse(png,png.name)['width']==320
    archive=tmp_path/'project.zip';archive.write_bytes(zip_fixture({'src/main.py':'answer = 42'}))
    assert parse(archive,archive.name)['segments'][0]['location']=='archive entry src/main.py'

def test_unsafe_and_unsupported_files(tmp_path):
    bad=tmp_path/'bad.zip';bad.write_bytes(zip_fixture({'../outside.py':'do not extract'}))
    with pytest.raises(ValueError,match='Unsafe'):parse(bad,bad.name)
    store=AssetStore(tmp_path)
    with pytest.raises(ValueError):store.create('script.exe',b'fake')
    with pytest.raises(ValueError):store.metadata('../../auth.sqlite3')
    store.executor.shutdown()

def test_formula_without_cached_value_is_not_invented(tmp_path):
    from openpyxl import Workbook
    book=Workbook();book.active.append(['Amount']);book.active.append(['=SUM(2,3)']);out=io.BytesIO();book.save(out)
    store=AssetStore(tmp_path);asset=store.create('formula.xlsx',out.getvalue());wait_ready(store,asset)
    with pytest.raises(ValueError,match='cached formula'):calculate_table(store,asset['id'],'Amount')
    store.executor.shutdown()

def test_authenticated_thumbnail_is_small_and_original_unchanged(tmp_path,monkeypatch):
    from PIL import Image
    store=AssetStore(tmp_path);monkeypatch.setattr('final_version_app.storage.assets.get_asset_store',lambda:store)
    image=Image.new('RGB',(1024,1024),'blue');buffer=io.BytesIO();image.save(buffer,format='PNG');original=buffer.getvalue()
    asset=store.create('blue.png',original);wait_ready(store,asset)
    client=TestClient(make_app(tmp_path));url=f"/api/files/{asset['id']}/thumbnail"
    assert client.get(url).status_code==401
    client.post('/api/auth/setup',json=CREDENTIALS)
    response=client.get(url);assert response.status_code==200
    with Image.open(io.BytesIO(response.content)) as thumbnail:assert thumbnail.size==(320,320)
    assert response.headers['content-type']=='image/jpeg'
    assert client.get(asset['download_url']).content==original
    assert client.get(url).content==response.content
    store.executor.shutdown()

def test_xlsx_aggregation_reports_actual_value_cells(tmp_path):
    from openpyxl import Workbook
    store=AssetStore(tmp_path);book=Workbook();book.active.title='Budget';book.active.append(['category','amount']);book.active.append(['A',120]);book.active.append(['B',80]);book.active.append(['A',40]);buffer=io.BytesIO();book.save(buffer)
    asset=store.create('budget.xlsx',buffer.getvalue());wait_ready(store,asset)
    result=calculate_table(store,asset['id'],'amount','sum',filter_column='category',filter_value='A')
    assert result['value']=='160'
    assert result['source_locations']==['sheet Budget, row 2, cell B2','sheet Budget, row 4, cell B4']
    duplicate=store.create('duplicate.csv',b'amount,amount\n1,2\n')
    with pytest.raises(ValueError,match='Duplicate'):calculate_table(store,duplicate['id'],'amount')
    store.executor.shutdown()
