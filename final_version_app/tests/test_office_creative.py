import io,json,time
from pathlib import Path
import pytest
from docx import Document
from openpyxl import load_workbook
from final_version_app.domain.office import document,spreadsheet
from final_version_app.storage.assets import AssetStore
from final_version_app.domain.skills import SkillLoader
from final_version_app.domain.creative_jobs import CreativeJobs

@pytest.fixture
def store(tmp_path):
    s=AssetStore(tmp_path);yield s;s.executor.shutdown()

def wav_bytes():
    import wave
    data=io.BytesIO()
    with wave.open(data,'wb') as audio:
        audio.setnchannels(1);audio.setsampwidth(2);audio.setframerate(8000);audio.writeframes(b'\0\0'*8000)
    return data.getvalue()

def mp3_bytes():
    return (bytes.fromhex('fffb9064')+b'\0'*413)*20

def test_word_cross_run_replacement_preserves_source_and_format(store):
    doc=Document();p=doc.add_paragraph();p.add_run('Budget: ');p.add_run('old ').bold=True;p.add_run('name');p.add_run(' / unchanged').italic=True
    doc.add_table(rows=1,cols=1).cell(0,0).text='old name'
    out=io.BytesIO();doc.save(out);original=store.create('source.docx',out.getvalue());before=store.path(original['id']).read_bytes()
    with pytest.raises(ValueError,match='matched 2'):document('edit',file_id=original['id'],replacements=[{'old':'old name','new':'new name'}],store=store)
    result=document('edit','revised.docx',original['id'],replacements=[{'old':'old name','new':'new name','expected_count':2}],store=store)
    changed=Document(store.path(result['asset']['id']));assert changed.paragraphs[0].text=='Budget: new name / unchanged'
    assert changed.paragraphs[0].runs[1].bold and changed.paragraphs[0].runs[-1].italic
    assert changed.tables[0].cell(0,0).text=='new name';assert store.path(original['id']).read_bytes()==before
    assert result['asset']['provenance']['parent_file_id']==original['id']

def test_excel_edit_formulas_chart_and_clear(store):
    made=spreadsheet('create',operations=[{'type':'write','start':'A1','rows':[['Item','Amount'],['A',120],['B',80],['Total','=SUM(B2:B3)']]},{'type':'format','range':'B2','number_format':'0.00'},{'type':'chart','range':'A1:B3'}],store=store)['asset']
    original=store.path(made['id']).read_bytes()
    edited=spreadsheet('edit','changed.xlsx',made['id'],operations=[{'type':'write','start':'B2','rows':[[150]]},{'type':'write','start':'A3','rows':[[None]]}],store=store)['asset']
    book=load_workbook(io.BytesIO(store.path(edited['id']).read_bytes()));assert book.active['B2'].value==150;assert book.active['B4'].value=='=SUM(B2:B3)';assert book.active['A3'].value is None
    assert book.active['B2'].number_format=='0.00';assert len(book.active._charts)==1;assert store.path(made['id']).read_bytes()==original
    seen=spreadsheet('inspect',file_id=edited['id'],cell_range='B4:B4',store=store);assert seen['cells'][0]['cached_value'] is None

def test_shared_skills_available_and_workspace_override(tmp_path,monkeypatch):
    common=tmp_path/'common';local=tmp_path/'workspace';(common/'sample').mkdir(parents=True);(local/'sample').mkdir(parents=True)
    (common/'sample/SKILL.md').write_text('---\nname: sample\ndescription: example\n---\ncommon')
    monkeypatch.setenv('AGENT_SHARED_SKILLS_DIR',str(common));assert 'common' in SkillLoader(tmp_path/'empty').load('sample')
    (local/'sample/SKILL.md').write_text('---\nname: sample\n---\nproject')
    assert 'project' in SkillLoader(local).load('sample')

def test_music_missing_config_and_unknown_recovery(store,monkeypatch):
    monkeypatch.delenv('ELEVENLABS_API_KEY',raising=False);jobs=CreativeJobs(store)
    with pytest.raises(ValueError,match='not configured'):jobs.start('music','test',prompt='music',seconds=30)
    now=time.time();job='job_'+'a'*32
    with store.connect() as db:db.execute('INSERT INTO creative_jobs VALUES (?,?,?,?,?,?,?,?)',(job,'recover',json.dumps({'kind':'music','prompt':'music','seconds':30}),'running',None,None,now,now))
    jobs.recover();assert jobs.status(job)['state']=='unknown';jobs.executor.shutdown()

def test_music_generate_persists_validated_mp3_asset(store,monkeypatch):
    class StreamResponse:
        status_code=200
        def __enter__(self):return self
        def __exit__(self,*args):return False
        def iter_bytes(self):
            data=mp3_bytes()
            yield data[:2000];yield data[2000:]
    class FakeClient:
        def __init__(self,*args,**kwargs):pass
        def __enter__(self):return self
        def __exit__(self,*args):return False
        def stream(self,*args,**kwargs):return StreamResponse()
    monkeypatch.setenv('ELEVENLABS_API_KEY','test-key')
    monkeypatch.setattr('final_version_app.domain.creative_jobs.httpx.Client',FakeClient)
    jobs=CreativeJobs(store)
    job=jobs.start('music','music-request',prompt='short piano loop',seconds=3)
    completed=jobs.check(job['id'],wait_seconds=30);assert completed['state']=='completed',completed
    asset=completed['result']['assets'][0];assert asset['status']=='ready' and asset['name'].endswith('.mp3')
    parsed=store.read(asset['id']);assert parsed['kind']=='media'
    assert parsed['media']['media_kind']=='audio' and parsed['media']['format']=='mp3'
    assert parsed['media']['sample_rate_hz']==44100 and parsed['media']['bitrate_bps']==128000
    jobs.executor.shutdown()

def test_office_job_idempotent_and_cancellation_before_execution(store,monkeypatch):
    from concurrent.futures import Future
    monkeypatch.setenv('AGENT_CREATIVE_URL','http://office:8094');jobs=CreativeJobs(store)
    monkeypatch.setattr(jobs.executor,'submit',lambda *args:Future())
    made=document('create',blocks=[{'type':'paragraph','text':'hello'}],store=store)['asset']
    first=jobs.start('office','same',file_id=made['id'],format='pdf');second=jobs.start('office','same',file_id=made['id'],format='pdf');assert first['id']==second['id']
    with pytest.raises(ValueError):jobs.start('office','same',file_id=made['id'],format='xlsx')
    assert jobs.cancel(first['id'])['state']=='cancelled';jobs._run(first['id']);assert jobs.status(first['id'])['state']=='cancelled';jobs.executor.shutdown()


def test_audio_upload_metadata_range_download_and_auth(store,monkeypatch,tmp_path):
    from fastapi.testclient import TestClient
    from final_version_app.tests.test_workspace_auth import make_app,CREDENTIALS
    monkeypatch.setattr('final_version_app.storage.assets.get_asset_store',lambda:store)
    data=wav_bytes()
    client=TestClient(make_app(tmp_path));assert client.get('/api/creative-jobs').status_code==401
    client.post('/api/auth/setup',json=CREDENTIALS)
    created=client.post('/api/files?name=sample.wav',content=data);assert created.status_code==201
    asset=created.json()
    for _ in range(100):
        if store.metadata(asset['id'])['status']!='queued' and store.metadata(asset['id'])['status']!='processing':break
        time.sleep(.05)
    assert store.metadata(asset['id'])['status']=='ready'
    parsed=store.read(asset['id']);assert parsed['kind']=='media'
    assert parsed['media']=={'media_kind':'audio','format':'wav','codec':'pcm','duration_seconds':1.0,'sample_rate_hz':8000,'channels':1,'sample_width_bits':16,'bitrate_bps':128000}
    response=client.get(asset['download_url'],headers={'Range':'bytes=0-31'});assert response.status_code==206
    assert response.content==data[:32];assert response.headers['content-type']=='audio/wav'
    client.post('/api/auth/logout');assert client.get(asset['download_url']).status_code==401
