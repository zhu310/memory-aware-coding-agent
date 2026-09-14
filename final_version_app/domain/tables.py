"""Deterministic numeric aggregation with explicit sheet/row evidence."""
import csv,io
from decimal import Decimal,InvalidOperation
from pathlib import Path
from final_version_app.document_parser import checked_zip,decode

def calculate_table(store,file_id,column,operation='sum',sheet='',filter_column='',filter_value=''):
    if operation not in {'sum','mean','min','max','count'}:raise ValueError('Unsupported operation')
    meta=store.metadata(file_id);path=store.path(file_id);suffix=Path(meta['name']).suffix.lower()
    rows=[];locations=[];warnings=[]
    if suffix=='.csv':
        reader=csv.DictReader(io.StringIO(decode(path.read_bytes())))
        columns=reader.fieldnames or []
        for i,row in enumerate(reader,2):
            if i>10001:raise ValueError('Table exceeds 10000 data rows; split it before calculating')
            rows.append(row);locations.append(f'CSV row {i}')
    elif suffix=='.xlsx':
        from openpyxl import load_workbook
        with checked_zip(path):pass
        source=path.read_bytes();book=load_workbook(io.BytesIO(source),read_only=True,data_only=True)
        formulas=load_workbook(io.BytesIO(source),read_only=True,data_only=False)
        try:
            table=book[sheet] if sheet else book.active
            formula_sheet=formulas[table.title]
            iterator=table.iter_rows(max_col=200)
            columns=[str(x.value or '') for x in next(iterator)]
            formula_rows=formula_sheet.iter_rows(min_row=2,max_col=200)
            for i,(row,formula_row) in enumerate(zip(iterator,formula_rows),2):
                if i>10001:raise ValueError('Table exceeds 10000 data rows')
                item={key:cell.value for key,cell in zip(columns,row) if key}
                for key,value,formula in zip(columns,row,formula_row):
                    if key==column and isinstance(formula.value,str) and formula.value.startswith('=') and value.value is None:
                        raise ValueError(f'{table.title}!{formula.coordinate} has no cached formula value; recalculate and save the workbook first')
                coordinate=next((cell.coordinate for key,cell in zip(columns,row) if key==column),'?')
                rows.append(item);locations.append(f'sheet {table.title}, row {i}, cell {coordinate}')
        finally:book.close();formulas.close()
    else:raise ValueError('table_calculate supports CSV and XLSX')
    named=[name for name in columns if name]
    if len(named)!=len(set(named)):raise ValueError('Duplicate column names are ambiguous; rename them before calculating')
    if column not in columns:raise ValueError(f'Unknown column; available columns: {columns}')
    if filter_column and filter_column not in columns:raise ValueError('Unknown filter column')
    selected=[(row,location) for row,location in zip(rows,locations) if not filter_column or str(row.get(filter_column,''))==filter_value]
    values=[];missing=0
    for row,location in selected:
        value=row.get(column)
        if value is None or value=='':missing+=1;continue
        if operation=='count':continue
        try:
            number=Decimal(str(value))
            if not number.is_finite():raise ValueError('Non-finite value')
            values.append(number)
        except (InvalidOperation,ValueError):raise ValueError(f'Non-numeric value at {location}, column {column}: {str(value)[:60]}')
    if operation=='count':answer=len(selected)
    elif not values:raise ValueError('No numeric values matched')
    elif operation=='sum':answer=sum(values,Decimal(0))
    elif operation=='mean':answer=sum(values,Decimal(0))/len(values)
    elif operation=='min':answer=min(values)
    else:answer=max(values)
    return {'file_id':file_id,'filename':meta['name'],'column':column,'operation':operation,'value':str(answer),'matched_rows':len(selected),'numeric_rows':len(values),'missing_values':missing,'source_locations':[x[1] for x in selected[:20]],'warnings':warnings}
