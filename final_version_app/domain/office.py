"""Bounded Office operations that always publish a new asset, never overwrite sources."""
from __future__ import annotations
import io,re,zipfile
from pathlib import Path
from final_version_app.storage.assets import get_asset_store
from final_version_app.document_parser import checked_zip


def _source(store,file_id,suffix):
    if Path(store.metadata(file_id)['name']).suffix.lower()!=suffix:raise ValueError('Expected '+suffix+' file')
    with checked_zip(store.path(file_id)):pass
    return store.path(file_id)


def _publish(store,name,suffix,data,parent,operation,warnings=None):
    if Path(name).name!=name or not name.lower().endswith(suffix):raise ValueError('Use a filename ending in '+suffix)
    asset=store.create(name,data,{'source':'office_tool','parent_file_id':parent or None,'operation':operation})
    return {'asset':asset,'source_preserved':True,'warnings':warnings or []}


def _paragraphs(container):
    yield from container.paragraphs
    for table in container.tables:
        seen=set()
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:continue
                seen.add(cell._tc);yield from _paragraphs(cell)


def _replace(paragraph,old,new):
    # Replace across runs, retaining the unaffected runs and formatting.
    runs=paragraph.runs;text=''.join(r.text for r in runs);starts=[];offset=0
    for r in runs:starts.append(offset);offset+=len(r.text)
    matches=list(re.finditer(re.escape(old),text))
    for match in reversed(matches):
        a,b=match.span();left=next(i for i,r in enumerate(runs) if starts[i]+len(r.text)>a)
        right=max(i for i in range(len(runs)) if starts[i]<b)
        if left==right:runs[left].text=runs[left].text[:a-starts[left]]+new+runs[left].text[b-starts[left]:]
        else:
            runs[left].text=runs[left].text[:a-starts[left]]+new
            for i in range(left+1,right):runs[i].text=''
            runs[right].text=runs[right].text[b-starts[right]:]
    return len(matches)


def document(action,name='document.docx',file_id='',blocks=None,replacements=None,store=None):
    from docx import Document
    from docx.shared import Pt,Cm,RGBColor
    from docx.oxml.ns import qn
    store=store or get_asset_store()
    if action not in {'inspect','create','edit'}:raise ValueError('action: inspect/create/edit')
    if action in {'inspect','edit'} and not file_id:raise ValueError('file_id is required')
    source=_source(store,file_id,'.docx') if file_id else None
    doc=Document(source) if source else Document()
    if action=='inspect':
        return {'file_id':file_id,'paragraphs':[{'index':i,'text':p.text[:4000],'style':p.style.name} for i,p in enumerate(doc.paragraphs[:150])],
                'tables':[[[c.text[:1000] for c in row.cells[:20]] for row in t.rows[:50]] for t in doc.tables[:15]],'inline_images':len(doc.inline_shapes),
                'guidance':'Inspection is bounded. Use file_read for additional text. Page layout requires office_render.'}
    blocks=blocks or [];replacements=replacements or []
    if len(blocks)>200 or len(replacements)>100:raise ValueError('Too many edits')
    if not source:
        normal=doc.styles['Normal'];normal.font.name='Noto Sans CJK SC';normal.font.size=Pt(10.5)
        normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');normal.paragraph_format.space_after=Pt(7)
        for section in doc.sections:section.top_margin=Cm(2);section.bottom_margin=Cm(2);section.left_margin=Cm(2.2);section.right_margin=Cm(2.2)
        for style_name in ['Title','Heading 1','Heading 2','Heading 3']:
            style=doc.styles[style_name];style.font.name='Noto Sans CJK SC';style.font.color.rgb=RGBColor.from_string('283C34')
            style.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    paragraphs=list(_paragraphs(doc))
    for section in doc.sections:
        for part in [section.header,section.footer]:paragraphs.extend(_paragraphs(part))
    changed=0
    for item in replacements:
        old=item.get('old','');new=item.get('new','')
        if not isinstance(old,str) or not old or not isinstance(new,str):raise ValueError('Each replacement requires nonempty old and string new')
        count=sum(''.join(r.text for r in p.runs).count(old) for p in paragraphs)
        if count!=item.get('expected_count',1):raise ValueError(f'Replacement matched {count}; provide the verified expected_count')
        for p in paragraphs:changed+=_replace(p,old,new)
    for block in blocks:
        kind=block.get('type','paragraph');text=str(block.get('text',''))
        if len(text)>20000:raise ValueError('Block text too large')
        if kind=='heading':doc.add_heading(text,level=max(0,min(3,int(block.get('level',1)))))
        elif kind in {'paragraph','bullet','numbered'}:doc.add_paragraph(text,style={'bullet':'List Bullet','numbered':'List Number'}.get(kind,'Normal'))
        elif kind=='table':
            rows=block.get('rows',[])
            if not rows or len(rows)>200 or not rows[0] or len(rows[0])>20 or any(len(r)!=len(rows[0]) for r in rows):raise ValueError('Use a rectangular table of at most 200 x 20')
            table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Light Shading Accent 1'
            for row in rows:
                for cell,value in zip(table.add_row().cells,row):cell.text=str(value)[:4000]
        elif kind=='image':
            image_id=block['file_id'];meta=store.metadata(image_id)
            if meta['media_type'] not in {'image/png','image/jpeg'}:raise ValueError('Document images require PNG/JPEG')
            doc.add_picture(str(store.path(image_id)),width=Cm(max(1,min(16,float(block.get('width_cm',12))))))
        elif kind=='page_break':doc.add_page_break()
        else:raise ValueError('Unknown block type: '+kind)
    output=io.BytesIO();doc.save(output)
    result=_publish(store,name,'.docx',output.getvalue(),file_id,action,['Complex text boxes, tracked changes and embedded objects require specialist review; layout is not verified until rendered.'])
    result['replacement_count']=changed;return result


def _cell_range(address):
    from openpyxl.utils.cell import range_boundaries
    try:a,b,c,d=range_boundaries(address)
    except Exception as exc:raise ValueError('Invalid cell range') from exc
    if None in (a,b,c,d) or min(a,b)<1 or c>200 or d>10000 or (c-a+1)*(d-b+1)>10000:raise ValueError('Range exceeds 10000 cells / 200 columns / 10000 rows')
    return a,b,c,d


def spreadsheet(action,name='workbook.xlsx',file_id='',operations=None,sheet='',cell_range='A1:T30',store=None):
    from openpyxl import Workbook,load_workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.workbook.properties import CalcProperties
    from openpyxl.chart import BarChart,LineChart,Reference
    store=store or get_asset_store()
    if action not in {'inspect','create','edit'}:raise ValueError('action: inspect/create/edit')
    if action in {'inspect','edit'} and not file_id:raise ValueError('file_id is required')
    source=_source(store,file_id,'.xlsx') if file_id else None
    if source:
        with zipfile.ZipFile(source) as z:
            if any('vbaProject' in n for n in z.namelist()):raise ValueError('Macro workbooks are not supported by this editor')
    book=load_workbook(io.BytesIO(source.read_bytes()),data_only=False) if source else Workbook()
    if action=='inspect':
        ws=book[sheet] if sheet else book.active;a,b,c,d=_cell_range(cell_range)
        cached=load_workbook(io.BytesIO(source.read_bytes()),data_only=True,read_only=True)
        result={'file_id':file_id,'sheets':book.sheetnames,'sheet':ws.title,'range':cell_range,'cells':[{'cell':cell.coordinate,'value':cell.value,'cached_value':cached[ws.title][cell.coordinate].value if cell.data_type=='f' else cell.value,'number_format':cell.number_format} for row in ws.iter_rows(min_row=b,max_row=d,min_col=a,max_col=c) for cell in row if cell.value is not None]}
        cached.close();book.close();return result
    operations=operations or []
    if len(operations)>100:raise ValueError('At most 100 operations')
    for op in operations:
        kind=op.get('type');target=op.get('sheet') or book.active.title
        if kind=='add_sheet':book.create_sheet(target);continue
        if target not in book.sheetnames:raise ValueError('Unknown sheet: '+target)
        ws=book[target]
        if kind=='rename_sheet':ws.title=op['title']
        elif kind=='write':
            rows=op.get('rows',[])
            if not rows or sum(len(r) for r in rows)>10000:raise ValueError('Write 1-10000 cells per operation')
            a,b,c,d=_cell_range(op.get('start','A1'))
            if b+len(rows)>10001 or a+max(map(len,rows))>201:raise ValueError('Write exceeds supported dimensions')
            for ri,row in enumerate(rows,b):
                for ci,value in enumerate(row,a):
                    if isinstance(value,(dict,list)):raise ValueError('Cell values must be scalar')
                    ws.cell(ri,ci).value=value
        elif kind in {'format','merge','filter'}:
            address=op['range'];a,b,c,d=_cell_range(address)
            if kind=='merge':ws.merge_cells(address)
            elif kind=='filter':ws.auto_filter.ref=address
            else:
                for row in ws.iter_rows(min_row=b,max_row=d,min_col=a,max_col=c):
                    for cell in row:
                        if 'number_format' in op:cell.number_format=op['number_format']
                        if 'bold' in op or 'color' in op:
                            from copy import copy
                            font=copy(cell.font)
                            if 'bold' in op:font.bold=bool(op['bold'])
                            if 'color' in op:font.color=op['color']
                            cell.font=font
                        if 'fill' in op:cell.fill=PatternFill('solid',fgColor=op['fill'])
                        if 'wrap' in op:cell.alignment=Alignment(wrap_text=bool(op['wrap']))
        elif kind=='freeze':_cell_range(op['cell']);ws.freeze_panes=op['cell']
        elif kind=='column_width':
            column=op['column']
            if not re.fullmatch('[A-Z]{1,2}',column):raise ValueError('Invalid column')
            ws.column_dimensions[column].width=max(3,min(80,float(op['width'])))
        elif kind=='chart':
            a,b,c,d=_cell_range(op['range']);chart=LineChart() if op.get('chart_type')=='line' else BarChart()
            if c<=a or d<=b:raise ValueError('Chart requires headers, categories and numeric values')
            chart.title=op.get('title','');chart.add_data(Reference(ws,min_col=a+1,max_col=c,min_row=b,max_row=d),titles_from_data=True)
            chart.set_categories(Reference(ws,min_col=a,min_row=b+1,max_row=d));anchor=op.get('anchor','G2');_cell_range(anchor);ws.add_chart(chart,anchor)
        else:raise ValueError('Unsupported spreadsheet operation: '+str(kind))
    book.calculation=CalcProperties(calcId=0,fullCalcOnLoad=True,forceFullCalc=True)
    output=io.BytesIO();book.save(output);book.close()
    return _publish(store,name,'.xlsx',output.getvalue(),file_id,action,['Formula expressions are saved but not evaluated by openpyxl. Use office_render(format=xlsx) to recalculate a copy. Advanced Excel objects may require specialist review.'])
